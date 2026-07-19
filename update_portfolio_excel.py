#!/usr/bin/env python3
"""Auto-update portfolio Excel workbook with latest market data."""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from portfolio_fetcher import (
    ALL_CODES,
    fetch_all_quotes,
    fetch_index_shanghai,
    fetch_limit_down_count,
)

DEFAULT_EXCEL = Path(__file__).resolve().parent / "持仓跟踪表_434050.xlsx"

# sheet -> (code_col, price_col, change_col or None, data_start_row, data_end_row)
SHEET_PRICE_MAP = {
    "当前持仓": ("C", "F", "K", 5, 15),
    "稳健版目标持仓": ("C", "F", None, 5, 13),
    "激进版目标持仓": ("C", "F", None, 5, 15),
    "每日监控": ("B", "E", None, 8, 15),
}


def _col_idx(letter: str) -> int:
    return ord(letter.upper()) - ord("A") + 1


def update_sheet_prices(ws, code_col: str, price_col: str, change_col: str | None,
                        start_row: int, end_row: int, quotes: dict) -> int:
    code_i = _col_idx(code_col)
    price_i = _col_idx(price_col)
    change_i = _col_idx(change_col) if change_col else None
    updated = 0

    for row in range(start_row, end_row + 1):
        code = ws.cell(row, code_i).value
        if not code or code in ("—", "-", ""):
            continue
        code = str(code).zfill(6)
        quote = quotes.get(code)
        if not quote:
            continue
        ws.cell(row, price_i, quote.price)
        if change_i and quote.change_pct is not None:
            ws.cell(row, change_i, quote.change_pct / 100)
            ws.cell(row, change_i).number_format = "0.00%"
        updated += 1
    return updated


def update_current_holdings_meta(ws, quotes: dict, index_quote, limit_down):
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    stock_dates = [q.trade_date for q in quotes.values()
                   if q.trade_date and q.source == "stock_zh_a_daily"]
    all_dates = [q.trade_date for q in quotes.values() if q.trade_date]
    data_date = max(stock_dates) if stock_dates else (max(all_dates) if all_dates else now[:10])
    ws["A2"] = f"数据更新：{now}  |  行情日期：{data_date}  |  账户总市值："
    ws["K2"] = "=SUM(G5:G15)"
    ws["K2"].number_format = "#,##0"


def update_daily_monitor_meta(ws, index_quote, limit_down):
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws["B3"] = now
    if index_quote:
        ws["B4"] = index_quote.price
        ws["B4"].number_format = "0.00"
    if limit_down is not None:
        ws["B5"] = limit_down


def update_data_log_sheet(wb, quotes, index_quote, limit_down):
    if "数据更新日志" not in wb.sheetnames:
        ws = wb.create_sheet("数据更新日志")
        ws.append(["更新时间", "代码", "名称", "现价", "涨跌幅%", "行情日期", "数据源"])
    else:
        ws = wb["数据更新日志"]

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    name_map = {}
    for sheet in ("当前持仓", "每日监控"):
        if sheet not in wb.sheetnames:
            continue
        sws = wb[sheet]
        code_col = 3 if sheet == "当前持仓" else 2
        name_col = 2 if sheet == "当前持仓" else 1
        start = 5 if sheet == "当前持仓" else 8
        end = 15
        for row in range(start, end + 1):
            code = sws.cell(row, code_col).value
            name = sws.cell(row, name_col).value
            if code and name:
                name_map[str(code).zfill(6)] = name

    for code in ALL_CODES:
        q = quotes.get(code)
        if not q:
            continue
        ws.append([
            now, code, name_map.get(code, ""),
            q.price,
            q.change_pct if q.change_pct is not None else "",
            q.trade_date or "",
            q.source,
        ])

    if index_quote:
        ws.append([
            now, "000001", "上证指数",
            index_quote.price, index_quote.change_pct,
            index_quote.trade_date, index_quote.source,
        ])
    if limit_down is not None:
        ws.append([now, "—", "跌停家数", limit_down, "", "", "stock_zh_a_spot_em"])


def run_update(excel_path: Path, save_as: Path | None = None) -> None:
    if not excel_path.exists():
        print(f"Excel not found: {excel_path}")
        print("Run: python3 generate_portfolio_excel.py")
        sys.exit(1)

    print(f"Loading {excel_path}...")
    wb = load_workbook(excel_path)

    print("Fetching market data (may take 1-2 minutes)...")
    quotes = fetch_all_quotes()
    missing = [c for c in ALL_CODES if c not in quotes]
    if missing:
        print(f"  [warn] missing quotes: {', '.join(missing)}")

    try:
        index_quote = fetch_index_shanghai()
        print(f"  上证指数: {index_quote.price} ({index_quote.change_pct}%)")
    except Exception as exc:  # noqa: BLE001
        print(f"  [warn] index: {exc}")
        index_quote = None

    limit_down = fetch_limit_down_count()
    if limit_down is not None:
        print(f"  跌停家数: {limit_down}")

    total_updated = 0
    for sheet_name, (code_col, price_col, change_col, start, end) in SHEET_PRICE_MAP.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        n = update_sheet_prices(ws, code_col, price_col, change_col, start, end, quotes)
        print(f"  Updated {sheet_name}: {n} rows")
        total_updated += n

    if "当前持仓" in wb.sheetnames:
        update_current_holdings_meta(wb["当前持仓"], quotes, index_quote, limit_down)

    if "每日监控" in wb.sheetnames:
        update_daily_monitor_meta(wb["每日监控"], index_quote, limit_down)

    update_data_log_sheet(wb, quotes, index_quote, limit_down)

    out = save_as or excel_path
    wb.save(out)
    print(f"\nDone. Updated {total_updated} price cells.")
    print(f"Saved: {out}")
    print(f"Time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")


def main():
    parser = argparse.ArgumentParser(description="Auto-update portfolio Excel with live quotes")
    parser.add_argument("-f", "--file", type=Path, default=DEFAULT_EXCEL, help="Excel file path")
    parser.add_argument("-o", "--output", type=Path, default=None, help="Save to different file")
    args = parser.parse_args()
    run_update(args.file, args.output)


if __name__ == "__main__":
    main()
